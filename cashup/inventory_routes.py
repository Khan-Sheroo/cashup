"""Inventory: live stock, invoices, requisitions, sales mix, and stock takes."""
from __future__ import annotations

import io
import json
import uuid
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from flask import (
    Blueprint, current_app, flash, redirect, render_template, request,
    send_file, send_from_directory, url_for,
)
from werkzeug.utils import secure_filename

from cashup import db
from cashup.inventory_import import (
    normalize_name, parse_invoice_csv, parse_item_sales_pdf, parse_purchase_order_pdf,
    parse_sales_csv, parse_stock_take_file, suggest_match, to_dec,
)
from cashup.inventory_util import last_movement_dates, on_hand_map, period_summary, explode_recipe
from cashup.category_rules import guess_category
from cashup.recipe_routes import _matching_category, category_suggestions
from cashup.recipe_util import INVOICE_QTY_UNIT_VALUES, INVOICE_QTY_UNITS, unit_conversion
from cashup.models import (
    MOVEMENT_KINDS, CostItem, Invoice, InvoiceLine, Location, PosItemAlias, Recipe, RecipeLine,
    Requisition, RequisitionLine, SalesImport, SalesLine, StockMovement, StockTake,
    StockTakeLine, Supplier, SupplierItemAlias,
)

inventory_bp = Blueprint('inventory', __name__, url_prefix='/inventory')

ZERO = Decimal('0')

ADJUST_KINDS = (
    ('receive', 'Receive (+)'),
    ('waste', 'Waste / Breakage (-)'),
    ('correction', 'Correction (+/-)'),
    ('opening', 'Opening Balance (+)'),
)


def _parse_date(raw, default=None):
    raw = (raw or '').strip()
    if not raw:
        return default
    try:
        return datetime.strptime(raw, '%Y-%m-%d').date()
    except ValueError:
        return default


def _parse_int(raw):
    try:
        return int(raw)
    except (TypeError, ValueError):
        return None


def _items_sorted(include_inactive=False):
    query = CostItem.query
    if not include_inactive:
        query = query.filter_by(active=True)
    items = query.all()
    items.sort(key=lambda i: ((i.category or '\uffff').lower(), i.name.lower()))
    return items


def _item_candidates():
    return {normalize_name(i.name): i.id for i in CostItem.query.filter_by(active=True).all()}


def _upload_dir(sub):
    path = Path(current_app.config['UPLOAD_FOLDER']) / sub
    path.mkdir(parents=True, exist_ok=True)
    return path


# ---------------------------------------------------------------------------
# Live stock
# ---------------------------------------------------------------------------

@inventory_bp.route('/')
def stock():
    locations = Location.ordered()
    selected = request.args.get('location', 'all')
    location = None
    if selected != 'all':
        location = Location.query.get(_parse_int(selected))
        if location is None:
            selected = 'all'

    per_location = {loc.id: on_hand_map(loc.id) for loc in locations}
    totals = on_hand_map(location.id if location else None)
    last_dates = last_movement_dates(location.id if location else None)

    selected_category = request.args.get('category', '')
    rows = []
    grand_value = ZERO
    categories: dict[str, int] = {}
    for item in _items_sorted(include_inactive=True):
        qty = totals.get(item.id, ZERO)
        if not item.active and qty == 0:
            continue
        categories[item.category or ''] = categories.get(item.category or '', 0) + 1
        if selected_category == '__none__' and item.category:
            continue
        if selected_category not in ('', '__none__') and item.category != selected_category:
            continue
        value = qty * item.cost_per_unit()
        grand_value += value
        par = Decimal(item.par_level) if item.par_level is not None else None
        count_qty = item.to_count_units(qty)
        rows.append({
            'item': item,
            'qty': qty,
            'count_qty': count_qty,
            'value': value,
            'par': par,
            'below_par': par is not None and count_qty < par,
            'last': last_dates.get(item.id),
            'by_location': {loc.id: item.to_count_units(per_location[loc.id].get(item.id, ZERO))
                            for loc in locations},
        })

    return render_template(
        'inventory/stock.html',
        rows=rows, locations=locations, location=location, selected=selected,
        grand_value=grand_value,
        category_counts=sorted(((c, n) for c, n in categories.items() if c), key=lambda x: x[0].lower()),
        uncategorised_count=categories.get('', 0),
        selected_category=selected_category,
        category_suggestions=category_suggestions(),
    )


@inventory_bp.route('/adjust', methods=['GET', 'POST'])
def adjust():
    locations = Location.ordered()
    items = _items_sorted()
    if request.method == 'POST':
        location = Location.query.get(_parse_int(request.form.get('location_id')))
        when = _parse_date(request.form.get('movement_date'), date.today())
        kind = request.form.get('kind', 'receive')
        note = (request.form.get('note') or '').strip() or None
        if location is None:
            flash('Choose a location', 'error')
            return render_template('inventory/adjust.html', locations=locations, items=items,
                                   adjust_kinds=ADJUST_KINDS, today=date.today())

        posted = 0
        item_ids = request.form.getlist('line_item_id')
        quantities = request.form.getlist('line_quantity')
        for idx, raw_item in enumerate(item_ids):
            item = CostItem.query.get(_parse_int(raw_item))
            qty = to_dec(quantities[idx] if idx < len(quantities) else None)
            if item is None or qty is None or qty == 0:
                continue
            base = item.to_base_units(qty)
            if kind == 'waste':
                movement_kind, base = 'waste', -abs(base)
            elif kind == 'opening':
                movement_kind, base = 'opening', abs(base)
            elif kind == 'receive':
                movement_kind, base = 'manual', abs(base)
            else:
                movement_kind = 'manual'
            db.session.add(StockMovement(
                item_id=item.id, location_id=location.id, movement_date=when,
                kind=movement_kind, quantity=base, unit_cost=item.cost_per_unit(), note=note,
            ))
            posted += 1
        if not posted:
            flash('Enter at least one item with a quantity', 'error')
            return render_template('inventory/adjust.html', locations=locations, items=items,
                                   adjust_kinds=ADJUST_KINDS, today=date.today())
        db.session.commit()
        flash(f'{posted} stock line{"s" if posted != 1 else ""} posted to {location.name}', 'success')
        return redirect(url_for('inventory.stock', location=location.id))

    return render_template('inventory/adjust.html', locations=locations, items=items,
                           adjust_kinds=ADJUST_KINDS, today=date.today())


@inventory_bp.route('/items/<int:item_id>/movements')
def item_movements(item_id):
    item = CostItem.query.get_or_404(item_id)
    movements = (item.movements
                 .order_by(StockMovement.movement_date.desc(), StockMovement.id.desc())
                 .limit(500).all())
    locations = Location.ordered()
    balances = {loc.id: on_hand_map(loc.id).get(item.id, ZERO) for loc in locations}
    return render_template('inventory/movements.html', item=item, movements=movements,
                           locations=locations, balances=balances)


@inventory_bp.route('/movements/<int:movement_id>/delete', methods=['POST'])
def delete_movement(movement_id):
    movement = StockMovement.query.get_or_404(movement_id)
    item_id = movement.item_id
    if movement.kind not in ('manual', 'waste', 'opening'):
        flash('Only manual entries can be deleted here — reverse the source document instead', 'error')
    else:
        db.session.delete(movement)
        db.session.commit()
        flash('Movement deleted', 'success')
    return redirect(url_for('inventory.item_movements', item_id=item_id))


# ---------------------------------------------------------------------------
# Invoices / purchase orders
# ---------------------------------------------------------------------------

@inventory_bp.route('/invoices')
def invoices():
    rows = Invoice.query.order_by(Invoice.invoice_date.desc(), Invoice.id.desc()).all()
    return render_template('inventory/invoices.html', invoices=rows)


@inventory_bp.route('/invoices/upload', methods=['POST'])
def upload_invoice():
    upload = request.files.get('file')
    store = Location.store()
    if not upload or not upload.filename:
        invoice = Invoice(invoice_date=date.today(), location_id=store.id if store else None)
        db.session.add(invoice)
        db.session.commit()
        return redirect(url_for('inventory.review_invoice', invoice_id=invoice.id))

    filename = upload.filename.lower()
    data = upload.read()
    try:
        if filename.endswith('.pdf'):
            parsed = parse_purchase_order_pdf(io.BytesIO(data))
        elif filename.endswith(('.csv', '.txt')):
            parsed = parse_invoice_csv(io.BytesIO(data))
        else:
            flash('Upload a PDF or CSV file', 'error')
            return redirect(url_for('inventory.invoices'))
    except Exception as exc:  # noqa: BLE001 - surface parse problems to the user
        flash(f'Could not read invoice: {exc}', 'error')
        return redirect(url_for('inventory.invoices'))

    stored = f'{uuid.uuid4().hex}_{secure_filename(upload.filename)}'
    (_upload_dir('invoices') / stored).write_bytes(data)

    supplier = None
    if parsed.get('supplier'):
        supplier = Supplier.query.filter(
            db.func.lower(Supplier.name) == parsed['supplier'].lower()
        ).first()
        if supplier is None:
            supplier = Supplier(name=parsed['supplier'])
            db.session.add(supplier)
            db.session.flush()

    invoice = Invoice(
        supplier_id=supplier.id if supplier else None,
        location_id=store.id if store else None,
        reference=parsed.get('reference'),
        invoice_date=parsed.get('date') or date.today(),
        original_filename=upload.filename,
        stored_filename=stored,
        stated_total=parsed.get('total'),
    )
    db.session.add(invoice)
    db.session.flush()

    candidates = _item_candidates()
    for order, row in enumerate(parsed['rows']):
        item_id, units, qty_unit = _match_invoice_line(supplier, row['description'], candidates)
        invoice.lines.append(InvoiceLine(
            description=row['description'][:255],
            quantity=row['quantity'],
            qty_unit=qty_unit,
            unit_price=row['unit_price'],
            tax_label=row.get('tax') or None,
            amount=row['amount'] or ZERO,
            item_id=item_id,
            units_per_line=units,
            sort_order=order,
        ))
    db.session.commit()

    if not parsed['rows']:
        flash('No line items were recognised — add them manually below', 'warning')
    else:
        matched = sum(1 for line in invoice.lines if line.item_id)
        flash(f'Read {len(invoice.lines)} lines, {matched} matched to items. Review then confirm.', 'success')
    return redirect(url_for('inventory.review_invoice', invoice_id=invoice.id))


def _match_invoice_line(supplier, description, candidates):
    key = normalize_name(description)
    alias = None
    if supplier is not None:
        alias = SupplierItemAlias.query.filter_by(supplier_id=supplier.id, description_key=key).first()
    if alias is None:
        alias = SupplierItemAlias.query.filter_by(description_key=key).first()
    if alias is not None:
        return alias.item_id, alias.units_per_line, alias.qty_unit or 'unit'
    item_id = suggest_match(description, candidates, cutoff=0.75)
    if item_id:
        item = CostItem.query.get(item_id)
        return item_id, item.pack_size if item else None, 'unit'
    return None, None, 'unit'


def _units_per_line(qty_unit, item, entered):
    """Recipe units per invoice quantity: converted for kg/g/L/ml, otherwise as entered."""
    if qty_unit != 'unit' and item is not None:
        factor = unit_conversion(qty_unit, item.unit)
        if factor is not None:
            return factor
    return entered


@inventory_bp.route('/invoices/<int:invoice_id>', methods=['GET', 'POST'])
def review_invoice(invoice_id):
    invoice = Invoice.query.get_or_404(invoice_id)
    if request.method == 'POST':
        if invoice.status == 'confirmed':
            flash('Invoice already confirmed — reverse it first to edit', 'error')
            return redirect(url_for('inventory.review_invoice', invoice_id=invoice.id))
        _save_invoice_form(invoice)
        action = request.form.get('action', 'save')
        errors = []
        if action == 'confirm':
            db.session.flush()
            errors = _confirm_invoice(invoice, request.form.get('update_cost') == 'on')
        db.session.commit()
        if errors:
            for err in errors[:8]:
                flash(err, 'error')
            return redirect(url_for('inventory.review_invoice', invoice_id=invoice.id))
        if action == 'confirm':
            flash(f'Invoice {invoice.reference or ""} confirmed — stock received'.replace('  ', ' '),
                  'success')
            return redirect(url_for('inventory.invoices'))
        else:
            flash('Invoice saved', 'success')
        return redirect(url_for('inventory.review_invoice', invoice_id=invoice.id))

    return render_template(
        'inventory/invoice_review.html',
        invoice=invoice,
        items=_items_sorted(),
        locations=Location.ordered(),
        suppliers=Supplier.query.order_by(Supplier.name).all(),
        qty_units=INVOICE_QTY_UNITS,
    )


def _save_invoice_form(invoice):
    form = request.form
    supplier_name = (form.get('supplier') or '').strip()
    if supplier_name:
        supplier = Supplier.query.filter(db.func.lower(Supplier.name) == supplier_name.lower()).first()
        if supplier is None:
            supplier = Supplier(name=supplier_name)
            db.session.add(supplier)
            db.session.flush()
        invoice.supplier_id = supplier.id
    else:
        invoice.supplier_id = None
    invoice.reference = (form.get('reference') or '').strip() or None
    invoice.invoice_date = _parse_date(form.get('invoice_date'), invoice.invoice_date or date.today())
    invoice.location_id = _parse_int(form.get('location_id')) or invoice.location_id

    existing = {line.id: line for line in invoice.lines}
    kept = set()
    order = 0
    for key in form.getlist('row_key'):
        desc = (form.get(f'desc_{key}') or '').strip()
        qty = to_dec(form.get(f'qty_{key}'))
        if not desc and qty is None:
            continue
        line = existing.get(_parse_int(key)) if not key.startswith('new') else None
        if line is None:
            line = InvoiceLine()
            invoice.lines.append(line)
        else:
            kept.add(line.id)
        line.description = desc[:255] or 'Item'
        line.quantity = qty or ZERO
        line.unit_price = to_dec(form.get(f'price_{key}'), ZERO)
        amount = to_dec(form.get(f'amount_{key}'))
        line.amount = amount if amount is not None else line.quantity * line.unit_price
        line.item_id = _parse_int(form.get(f'item_{key}'))
        qty_unit = form.get(f'qtyunit_{key}', 'unit')
        line.qty_unit = qty_unit if qty_unit in INVOICE_QTY_UNIT_VALUES else 'unit'
        item = CostItem.query.get(line.item_id) if line.item_id else None
        line.units_per_line = _units_per_line(line.qty_unit, item, to_dec(form.get(f'units_{key}')))
        line.skip = form.get(f'skip_{key}') == 'on'
        line.sort_order = order
        order += 1
    for line_id, line in existing.items():
        if line_id not in kept:
            invoice.lines.remove(line)


# Invoice quantity unit -> (recipe unit, pack size) for items created from an invoice line
NEW_ITEM_UNITS = {
    'unit': ('each', Decimal('1')),
    'kg': ('g', Decimal('1000')),
    'g': ('g', Decimal('1')),
    'L': ('ml', Decimal('1000')),
    'ml': ('ml', Decimal('1')),
}


def _line_price_per_qty(line) -> Decimal:
    """Price per invoice quantity, taken from the line amount so re-stating qty never inflates cost."""
    qty = Decimal(line.quantity or 0)
    if line.amount is not None and qty > 0:
        return Decimal(line.amount) / qty
    return Decimal(line.unit_price or 0)


def _line_cost_per_base_unit(line) -> Decimal:
    units = Decimal(line.units_per_line or 0)
    return _line_price_per_qty(line) / units if units > 0 else ZERO


def _item_for_unmatched_line(line):
    """Existing item with the same name, or a new item priced from the invoice line."""
    name = line.description.strip()[:200] or 'Item'
    item = CostItem.query.filter(db.func.lower(CostItem.name) == name.lower()).first()
    if item is not None:
        return item, False
    unit, pack_size = NEW_ITEM_UNITS.get(line.qty_unit or 'unit', NEW_ITEM_UNITS['unit'])
    item = CostItem(
        name=name,
        unit=unit,
        pack_size=pack_size,
        cost_price=_line_price_per_qty(line) * pack_size / (
            unit_conversion(line.qty_unit, unit) or Decimal('1')),
        count_unit=None if unit == 'each' else line.qty_unit,
        count_factor=None,
        category=_matching_category(guess_category(name)),
        active=True,
    )
    db.session.add(item)
    db.session.flush()
    return item, True


def _confirm_invoice(invoice, update_cost):
    errors = []
    if not invoice.location_id:
        errors.append('Choose the location receiving this stock')
    active_lines = [l for l in invoice.lines if not l.skip]
    if not active_lines:
        errors.append('No lines to receive')
    for line in active_lines:
        if line.item_id and (not line.units_per_line or line.units_per_line <= 0):
            errors.append(f'"{line.description}": enter units received per quantity')
    if errors:
        return errors

    created = []
    for line in active_lines:
        if not line.item_id:
            item, is_new = _item_for_unmatched_line(line)
            line.item_id = item.id
            line.units_per_line = _units_per_line(line.qty_unit, item, None) or Decimal('1')
            if is_new:
                created.append(item.name)
    if created:
        flash(f'Added {len(created)} new item{"s" if len(created) != 1 else ""} to inventory: '
              f'{", ".join(created[:10])}{"…" if len(created) > 10 else ""}. '
              f'Set their category and count unit on the Items page.', 'info')

    supplier = Supplier.query.get(invoice.supplier_id) if invoice.supplier_id else None
    note = f'{supplier.name if supplier else "Invoice"} {invoice.reference or ""}'.strip()
    for line in active_lines:
        base = line.base_quantity()
        unit_cost = _line_cost_per_base_unit(line)
        db.session.add(StockMovement(
            item_id=line.item_id, location_id=invoice.location_id,
            movement_date=invoice.invoice_date, kind='purchase', quantity=base,
            unit_cost=unit_cost, invoice_id=invoice.id, note=note,
        ))
        key = normalize_name(line.description)
        alias = SupplierItemAlias.query.filter_by(
            supplier_id=invoice.supplier_id, description_key=key
        ).first()
        if alias is None:
            alias = SupplierItemAlias(supplier_id=invoice.supplier_id, description_key=key)
            db.session.add(alias)
        alias.item_id = line.item_id
        alias.units_per_line = line.units_per_line
        alias.qty_unit = line.qty_unit or 'unit'
        if update_cost and unit_cost:
            item = CostItem.query.get(line.item_id)
            item.cost_price = unit_cost * Decimal(item.pack_size or 1)
    invoice.status = 'confirmed'
    invoice.confirmed_at = datetime.utcnow()
    return []


@inventory_bp.route('/invoices/<int:invoice_id>/reverse', methods=['POST'])
def reverse_invoice(invoice_id):
    invoice = Invoice.query.get_or_404(invoice_id)
    StockMovement.query.filter_by(invoice_id=invoice.id).delete()
    invoice.status = 'draft'
    invoice.confirmed_at = None
    db.session.commit()
    flash('Invoice reversed — stock removed, now editable', 'success')
    return redirect(url_for('inventory.review_invoice', invoice_id=invoice.id))


@inventory_bp.route('/invoices/<int:invoice_id>/delete', methods=['POST'])
def delete_invoice(invoice_id):
    invoice = Invoice.query.get_or_404(invoice_id)
    StockMovement.query.filter_by(invoice_id=invoice.id).delete()
    if invoice.stored_filename:
        path = _upload_dir('invoices') / invoice.stored_filename
        if path.exists():
            path.unlink()
    db.session.delete(invoice)
    db.session.commit()
    flash('Invoice deleted', 'success')
    return redirect(url_for('inventory.invoices'))


@inventory_bp.route('/invoices/<int:invoice_id>/file')
def invoice_file(invoice_id):
    invoice = Invoice.query.get_or_404(invoice_id)
    if not invoice.stored_filename:
        flash('No file attached', 'error')
        return redirect(url_for('inventory.review_invoice', invoice_id=invoice.id))
    return send_from_directory(_upload_dir('invoices'), invoice.stored_filename,
                               download_name=invoice.original_filename)


# ---------------------------------------------------------------------------
# Requisitions (store room -> outlet)
# ---------------------------------------------------------------------------

@inventory_bp.route('/requisitions')
def requisitions():
    rows = Requisition.query.order_by(Requisition.req_date.desc(), Requisition.id.desc()).all()
    return render_template('inventory/requisitions.html', requisitions=rows)


@inventory_bp.route('/requisitions/new', methods=['GET', 'POST'])
def new_requisition():
    locations = Location.ordered()
    store = Location.store()
    items = _items_sorted()
    store_stock = on_hand_map(store.id) if store else {}

    def render():
        return render_template('inventory/requisition_form.html', locations=locations,
                               store=store, items=items, store_stock=store_stock,
                               today=date.today())

    if request.method == 'POST':
        from_loc = Location.query.get(_parse_int(request.form.get('from_location_id')))
        to_loc = Location.query.get(_parse_int(request.form.get('to_location_id')))
        if from_loc is None or to_loc is None or from_loc.id == to_loc.id:
            flash('Choose two different locations', 'error')
            return render()
        req = Requisition(
            reference=(request.form.get('reference') or '').strip() or None,
            req_date=_parse_date(request.form.get('req_date'), date.today()),
            from_location_id=from_loc.id, to_location_id=to_loc.id,
            note=(request.form.get('note') or '').strip() or None,
        )
        item_ids = request.form.getlist('line_item_id')
        quantities = request.form.getlist('line_quantity')
        for idx, raw_item in enumerate(item_ids):
            item = CostItem.query.get(_parse_int(raw_item))
            qty = to_dec(quantities[idx] if idx < len(quantities) else None)
            if item is None or qty is None or qty <= 0:
                continue
            req.lines.append(RequisitionLine(item_id=item.id, quantity=qty))
        if not req.lines:
            flash('Add at least one item with a quantity', 'error')
            return render()
        db.session.add(req)
        db.session.flush()
        note = f'Req {req.reference or req.id}: {from_loc.name} → {to_loc.name}'
        for line in req.lines:
            item = CostItem.query.get(line.item_id)
            base = item.to_base_units(line.quantity)
            for loc_id, kind, qty in ((from_loc.id, 'transfer_out', -base),
                                      (to_loc.id, 'transfer_in', base)):
                db.session.add(StockMovement(
                    item_id=item.id, location_id=loc_id, movement_date=req.req_date,
                    kind=kind, quantity=qty, unit_cost=item.cost_per_unit(),
                    requisition_id=req.id, note=note,
                ))
        db.session.commit()
        flash(f'Requisition saved — {len(req.lines)} items moved to {to_loc.name}', 'success')
        return redirect(url_for('inventory.view_requisition', req_id=req.id))

    return render()


@inventory_bp.route('/requisitions/<int:req_id>')
def view_requisition(req_id):
    req = Requisition.query.get_or_404(req_id)
    return render_template('inventory/requisition_view.html', req=req)


@inventory_bp.route('/requisitions/<int:req_id>/delete', methods=['POST'])
def delete_requisition(req_id):
    req = Requisition.query.get_or_404(req_id)
    StockMovement.query.filter_by(requisition_id=req.id).delete()
    db.session.delete(req)
    db.session.commit()
    flash('Requisition deleted and stock moved back', 'success')
    return redirect(url_for('inventory.requisitions'))


# ---------------------------------------------------------------------------
# Sales mix
# ---------------------------------------------------------------------------

@inventory_bp.route('/sales')
def sales():
    rows = SalesImport.query.order_by(SalesImport.end_date.desc(), SalesImport.id.desc()).all()
    return render_template('inventory/sales.html', imports=rows)


@inventory_bp.route('/sales/upload', methods=['POST'])
def upload_sales():
    upload = request.files.get('file')
    if not upload or not upload.filename:
        flash('Choose a sales mix file', 'error')
        return redirect(url_for('inventory.sales'))
    filename = upload.filename.lower()
    try:
        if filename.endswith('.pdf'):
            parsed = parse_item_sales_pdf(upload)
        elif filename.endswith(('.csv', '.txt')):
            parsed = parse_sales_csv(upload)
        else:
            flash('Upload the POS Item Sales PDF or a CSV', 'error')
            return redirect(url_for('inventory.sales'))
    except Exception as exc:  # noqa: BLE001
        flash(f'Could not read sales mix: {exc}', 'error')
        return redirect(url_for('inventory.sales'))

    if not parsed['rows']:
        flash('No sales rows were recognised in that file', 'error')
        return redirect(url_for('inventory.sales'))

    start = parsed['start'] or _parse_date(request.form.get('start_date'), date.today())
    end = parsed['end'] or _parse_date(request.form.get('end_date'), start)
    outlets = [loc for loc in Location.ordered() if loc.kind != 'store']
    sales_import = SalesImport(
        start_date=start, end_date=end, original_filename=upload.filename,
        default_location_id=outlets[0].id if outlets else None,
    )
    db.session.add(sales_import)
    db.session.flush()

    recipe_candidates = {normalize_name(r.name): r.id for r in Recipe.query.all()}
    item_candidates = _item_candidates()
    aliases = {a.pos_name_key: a for a in PosItemAlias.query.all()}
    for order, row in enumerate(parsed['rows']):
        key = normalize_name(row['name'])
        alias = aliases.get(key)
        recipe_id = item_id = None
        ignore = False
        if alias is not None:
            recipe_id, item_id, ignore = alias.recipe_id, alias.item_id, alias.ignore
        else:
            recipe_id = suggest_match(row['name'], recipe_candidates, cutoff=0.85)
            if recipe_id is None:
                item_id = suggest_match(row['name'], item_candidates, cutoff=0.85)
        sales_import.lines.append(SalesLine(
            pos_name=row['name'][:255], quantity=row['quantity'], gross=row['gross'],
            discount=row['discount'], net=row['net'], recipe_id=recipe_id, item_id=item_id,
            ignore=ignore, sort_order=order,
        ))
    db.session.commit()
    mapped = sum(1 for l in sales_import.lines if l.is_mapped)
    flash(f'Read {len(sales_import.lines)} items ({mapped} already mapped). '
          f'Map the rest to recipes or items, then confirm.', 'success')
    return redirect(url_for('inventory.review_sales', import_id=sales_import.id))


@inventory_bp.route('/sales/<int:import_id>', methods=['GET', 'POST'])
def review_sales(import_id):
    sales_import = SalesImport.query.get_or_404(import_id)
    if request.method == 'POST':
        if sales_import.status == 'confirmed':
            flash('Already confirmed — reverse it first to edit', 'error')
            return redirect(url_for('inventory.review_sales', import_id=sales_import.id))
        _save_sales_form(sales_import)
        action = request.form.get('action', 'save')
        if action == 'confirm':
            errors = _confirm_sales(sales_import)
            if errors:
                db.session.commit()
                for err in errors[:8]:
                    flash(err, 'error')
                return redirect(url_for('inventory.review_sales', import_id=sales_import.id))
        db.session.commit()
        flash('Sales confirmed — stock usage posted' if action == 'confirm' else 'Mappings saved',
              'success')
        return redirect(url_for('inventory.review_sales', import_id=sales_import.id))

    recipes = Recipe.query.order_by(Recipe.name).all()
    lines = sorted(sales_import.lines, key=lambda l: (l.is_mapped, l.sort_order))
    return render_template(
        'inventory/sales_review.html',
        sales_import=sales_import, lines=lines, recipes=recipes, items=_items_sorted(),
        locations=[loc for loc in Location.ordered()],
        unmapped=sum(1 for l in sales_import.lines if not l.is_mapped),
    )


def _remember_pos_mapping(line):
    key = normalize_name(line.pos_name)
    alias = PosItemAlias.query.filter_by(pos_name_key=key).first()
    if alias is None:
        alias = PosItemAlias(pos_name_key=key)
        db.session.add(alias)
    alias.recipe_id = line.recipe_id
    alias.item_id = line.item_id
    alias.ignore = line.ignore


def _save_sales_form(sales_import):
    form = request.form
    sales_import.start_date = _parse_date(form.get('start_date'), sales_import.start_date)
    sales_import.end_date = _parse_date(form.get('end_date'), sales_import.end_date)
    sales_import.default_location_id = _parse_int(form.get('default_location_id'))
    for line in sales_import.lines:
        mode = form.get(f'mode_{line.id}')
        if mode is not None:
            pick = form.get(f'pick_{line.id}') or ''
            value = {'newitem': '', 'newrecipe': 'newrecipe', 'ignore': 'ignore'}.get(mode, pick)
        else:
            value = form.get(f'map_{line.id}')
        if value is None:
            continue
        line.recipe_id = None
        line.item_id = None
        line.ignore = False
        if value == 'ignore':
            line.ignore = True
        elif value == 'newrecipe':
            line.recipe_id = _empty_recipe_for_sale(line, sales_import.default_location_id).id
        elif value.startswith('item:'):
            line.item_id = _parse_int(value[5:])
        elif value.startswith('recipe:'):
            line.recipe_id = _parse_int(value[7:])
        if line.is_mapped:
            _remember_pos_mapping(line)


def _empty_recipe_for_sale(line, location_id):
    """Existing recipe with the POS name, or a new one with no ingredients yet."""
    name = line.pos_name.strip()[:200] or 'POS Item'
    recipe = Recipe.query.filter(db.func.lower(Recipe.name) == name.lower()).first()
    if recipe is not None:
        return recipe
    sale_price = None
    if line.gross and line.quantity:
        sale_price = (Decimal(line.gross) / Decimal(line.quantity)).quantize(Decimal('0.01'))
    recipe = Recipe(name=name, sale_price=sale_price, location_id=location_id,
                    notes='Created from sales mix — add ingredients')
    db.session.add(recipe)
    db.session.flush()
    return recipe


def _item_for_unmapped_sale(line):
    """Existing item with the POS name, or a new 1-each item. Returns (item, created)."""
    name = line.pos_name.strip()[:200] or 'POS Item'
    item = CostItem.query.filter(db.func.lower(CostItem.name) == name.lower()).first()
    if item is not None:
        return item, False
    item = CostItem(name=name, unit='each', pack_size=Decimal('1'), cost_price=ZERO,
                    category=_matching_category(guess_category(name, menu_item=True)) or 'From Sales Mix',
                    active=True)
    db.session.add(item)
    db.session.flush()
    return item, True


def _confirm_sales(sales_import):
    errors = []
    default_location = sales_import.default_location_id
    unmapped = [l for l in sales_import.lines if not l.is_mapped]
    item_lines = [l for l in sales_import.lines if l.item_id and not l.ignore]
    if (unmapped or item_lines) and not default_location:
        errors.append('Choose a Default Location — items sold as-is and new items are '
                      'taken from that location')
    for line in sales_import.lines:
        if line.ignore or not line.recipe_id:
            continue
        recipe = Recipe.query.get(line.recipe_id)
        if recipe.lines and not (recipe.location_id or default_location):
            errors.append(f'Recipe "{recipe.name}" has no location — set one or choose a default')
    if errors:
        return errors

    empty = sorted({Recipe.query.get(l.recipe_id).name for l in sales_import.lines
                    if l.recipe_id and not l.ignore and not Recipe.query.get(l.recipe_id).lines})
    if empty:
        flash(f'{len(empty)} recipe{"s have" if len(empty) != 1 else " has"} no ingredients yet, so '
              f'no stock was used for: {", ".join(empty[:10])}{"…" if len(empty) > 10 else ""}. '
              f'Add ingredients, then Reverse and Confirm this sales mix again to include them.',
              'warning')

    created = []
    for line in unmapped:
        item, is_new = _item_for_unmapped_sale(line)
        line.item_id = item.id
        _remember_pos_mapping(line)
        if is_new:
            created.append(item.name)
    if created:
        flash(f'Added {len(created)} new item{"s" if len(created) != 1 else ""} from the sales mix '
              f'(auto-categorised where recognised, otherwise "From Sales Mix"): '
              f'{", ".join(created[:10])}'
              f'{"…" if len(created) > 10 else ""}. Set their cost and count unit on the Items page.',
              'info')

    usage: dict[tuple[int, int], Decimal] = {}

    def add_usage(item_id, location_id, qty):
        usage[(item_id, location_id)] = usage.get((item_id, location_id), ZERO) + qty

    for line in sales_import.lines:
        if line.ignore:
            continue
        if line.recipe_id:
            recipe = Recipe.query.get(line.recipe_id)
            location_id = recipe.location_id or default_location
            for item_id, qty in explode_recipe(recipe, line.quantity).items():
                add_usage(item_id, location_id, qty)
        elif line.item_id:
            item = CostItem.query.get(line.item_id)
            add_usage(item.id, default_location, Decimal(line.quantity or 0) * Decimal(item.pack_size or 1))

    note = f'Sales {sales_import.start_date:%d/%m} – {sales_import.end_date:%d/%m/%Y}'
    items = {i.id: i for i in CostItem.query.filter(
        CostItem.id.in_({k[0] for k in usage})).all()} if usage else {}
    for (item_id, location_id), qty in usage.items():
        if qty == 0:
            continue
        db.session.add(StockMovement(
            item_id=item_id, location_id=location_id, movement_date=sales_import.end_date,
            kind='sale', quantity=-qty, unit_cost=items[item_id].cost_per_unit(),
            sales_import_id=sales_import.id, note=note,
        ))
    sales_import.status = 'confirmed'
    sales_import.confirmed_at = datetime.utcnow()
    return []


@inventory_bp.route('/sales/<int:import_id>/reverse', methods=['POST'])
def reverse_sales(import_id):
    sales_import = SalesImport.query.get_or_404(import_id)
    StockMovement.query.filter_by(sales_import_id=sales_import.id).delete()
    sales_import.status = 'draft'
    sales_import.confirmed_at = None
    db.session.commit()
    flash('Sales usage reversed — now editable', 'success')
    return redirect(url_for('inventory.review_sales', import_id=sales_import.id))


@inventory_bp.route('/sales/<int:import_id>/delete', methods=['POST'])
def delete_sales(import_id):
    sales_import = SalesImport.query.get_or_404(import_id)
    StockMovement.query.filter_by(sales_import_id=sales_import.id).delete()
    db.session.delete(sales_import)
    db.session.commit()
    flash('Sales import deleted', 'success')
    return redirect(url_for('inventory.sales'))


# ---------------------------------------------------------------------------
# Stock takes
# ---------------------------------------------------------------------------

@inventory_bp.route('/stocktakes')
def stocktakes():
    rows = StockTake.query.order_by(StockTake.count_date.desc(), StockTake.id.desc()).all()
    return render_template('inventory/stocktakes.html', stocktakes=rows,
                           locations=Location.ordered(), today=date.today())


@inventory_bp.route('/stocktakes/new', methods=['POST'])
def new_stocktake():
    location = Location.query.get(_parse_int(request.form.get('location_id')))
    if location is None:
        flash('Choose a location', 'error')
        return redirect(url_for('inventory.stocktakes'))
    take = StockTake(
        location_id=location.id,
        count_date=_parse_date(request.form.get('count_date'), date.today()),
        note=(request.form.get('note') or '').strip() or None,
    )
    db.session.add(take)
    db.session.commit()
    return redirect(url_for('inventory.count_stocktake', take_id=take.id))


def _count_items(take, show_all=False, categories=None):
    """Items shown on a count sheet: anything with history at the location, or all active.

    categories: optional list of category names to keep ('' = Uncategorised).
    """
    counted_ids = {line.item_id for line in take.lines}
    if show_all:
        items = _items_sorted()
        extra = [i for i in CostItem.query.filter(CostItem.id.in_(counted_ids)).all()
                 if i not in items] if counted_ids else []
        items.extend(extra)
    else:
        moved_ids = {row[0] for row in db.session.query(StockMovement.item_id)
                     .filter(StockMovement.location_id == take.location_id).distinct()}
        ids = moved_ids | counted_ids
        items = CostItem.query.filter(CostItem.id.in_(ids)).all() if ids else []
    if categories:
        wanted = set(categories)
        items = [i for i in items if (i.category or '') in wanted]
    items.sort(key=lambda i: ((i.category or '\uffff').lower(), i.name.lower()))
    return items


def _count_sheet_args():
    """show_all / categories from the query string, shared by the print and Excel sheets."""
    show_all = request.args.get('only') != '1'
    categories = [c for c in request.args.getlist('cat')] or None
    return show_all, categories


@inventory_bp.route('/stocktakes/<int:take_id>/count-sheet')
def print_count_sheet(take_id):
    take = StockTake.query.get_or_404(take_id)
    show_all, categories = _count_sheet_args()
    items = _count_items(take, show_all, categories)
    counts = {line.item_id: line.counted for line in take.lines}
    return render_template('inventory/count_sheet_print.html', items=items,
                           location_name=take.location.name, count_date=take.count_date,
                           back_url=url_for('inventory.count_stocktake', take_id=take.id),
                           counts=counts, with_counts=request.args.get('counts') == '1',
                           categories=categories)


@inventory_bp.route('/count-sheet')
def print_blank_count_sheet():
    """Full categorised count sheet for a location, without opening a stock take first."""
    location = Location.query.get(
        _parse_int(request.args.get('location') or request.args.get('location_id')))
    categories = request.args.getlist('cat') or None
    items = _items_sorted()
    if categories:
        items = [i for i in items if (i.category or '') in set(categories)]
    return render_template('inventory/count_sheet_print.html', items=items,
                           location_name=location.name if location else 'All Locations',
                           count_date=None, back_url=url_for('inventory.stocktakes'),
                           counts={}, with_counts=False, categories=categories)


@inventory_bp.route('/stocktakes/<int:take_id>', methods=['GET', 'POST'])
def count_stocktake(take_id):
    take = StockTake.query.get_or_404(take_id)
    show_all = request.args.get('only') != '1'
    if request.method == 'POST':
        if take.status == 'confirmed':
            flash('Stock take is confirmed — reverse it first to edit', 'error')
            return redirect(url_for('inventory.count_stocktake', take_id=take.id))
        take.count_date = _parse_date(request.form.get('count_date'), take.count_date)
        take.note = (request.form.get('note') or '').strip() or None
        existing = {line.item_id: line for line in take.lines}
        for key, raw in request.form.items():
            if not key.startswith('count_'):
                continue
            item_id = _parse_int(key[6:])
            if item_id is None:
                continue
            value = to_dec(raw)
            line = existing.get(item_id)
            if value is None:
                if line is not None:
                    take.lines.remove(line)
                continue
            if line is None:
                line = StockTakeLine(item_id=item_id)
                take.lines.append(line)
            line.counted = value
        action = request.form.get('action', 'save')
        if action == 'confirm':
            if not take.lines:
                flash('Enter at least one count before confirming', 'error')
                db.session.commit()
                return redirect(url_for('inventory.count_stocktake', take_id=take.id))
            db.session.flush()
            _confirm_stocktake(take)
            db.session.commit()
            flash('Stock take confirmed — stock levels updated to the count', 'success')
            return redirect(url_for('inventory.variance', take_id=take.id))
        db.session.commit()
        flash('Counts saved', 'success')
        return redirect(url_for('inventory.count_stocktake', take_id=take.id,
                                only='1' if request.form.get('show_all') == '0' else None))

    counts = {line.item_id: line.counted for line in take.lines}
    sheet_categories = sorted({i.category or '' for i in _count_items(take, show_all)},
                              key=lambda c: (c == '', c.lower()))
    return render_template('inventory/stocktake_count.html', take=take,
                           items=_count_items(take, show_all), counts=counts, show_all=show_all,
                           sheet_categories=sheet_categories)


def _confirm_stocktake(take):
    summary = period_summary(take)
    rows = {row['item'].id: row for row in summary['rows']}
    for line in take.lines:
        row = rows.get(line.item_id)
        item = CostItem.query.get(line.item_id)
        expected = row['expected'] if row else ZERO
        line.expected_base = expected
        line.unit_cost = item.cost_per_unit()
        diff = item.to_base_units(line.counted) - expected
        if diff != 0:
            db.session.add(StockMovement(
                item_id=item.id, location_id=take.location_id, movement_date=take.count_date,
                kind='stocktake_adjust', quantity=diff, unit_cost=line.unit_cost,
                stock_take_id=take.id, note=f'Stock take {take.count_date:%d/%m/%Y}',
            ))
    take.status = 'confirmed'
    take.confirmed_at = datetime.utcnow()


@inventory_bp.route('/stocktakes/<int:take_id>/upload', methods=['POST'])
def upload_stocktake(take_id):
    take = StockTake.query.get_or_404(take_id)
    if take.status == 'confirmed':
        flash('Stock take is confirmed — reverse it first', 'error')
        return redirect(url_for('inventory.count_stocktake', take_id=take.id))
    upload = request.files.get('file')
    if not upload or not upload.filename:
        flash('Choose a stock take file', 'error')
        return redirect(url_for('inventory.count_stocktake', take_id=take.id))
    try:
        rows = parse_stock_take_file(upload)
    except Exception as exc:  # noqa: BLE001
        flash(f'Could not read stock take: {exc}', 'error')
        return redirect(url_for('inventory.count_stocktake', take_id=take.id))

    candidates = {normalize_name(i.name): i.id for i in CostItem.query.all()}
    existing = {line.item_id: line for line in take.lines}
    unmatched = []
    applied = 0
    for row in rows:
        item_id = row['item_id'] if row['item_id'] and CostItem.query.get(row['item_id']) else None
        if item_id is None:
            item_id = suggest_match(row['description'], candidates, cutoff=0.85)
        if item_id is None:
            unmatched.append({'description': row['description'], 'count': str(row['count'])})
            continue
        line = existing.get(item_id)
        if line is None:
            line = StockTakeLine(item_id=item_id)
            take.lines.append(line)
            existing[item_id] = line
        line.counted = row['count']
        applied += 1
    take.unmatched_json = json.dumps(unmatched) if unmatched else None
    db.session.commit()
    msg = f'Applied {applied} counts from {upload.filename}'
    if unmatched:
        msg += f' — {len(unmatched)} rows did not match an item (listed below)'
    flash(msg, 'warning' if unmatched else 'success')
    return redirect(url_for('inventory.count_stocktake', take_id=take.id))


@inventory_bp.route('/stocktakes/<int:take_id>/template.xlsx')
def stocktake_template(take_id):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    take = StockTake.query.get_or_404(take_id)
    show_all, categories = _count_sheet_args()
    counts = {line.item_id: line.counted for line in take.lines}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Count'
    ws.append([f'{take.location.name} Stock Take', None, f'Count date: {take.count_date:%d/%m/%Y}'])
    ws.cell(1, 1).font = Font(bold=True, size=14)
    ws.append(['Counted by: ____________________', None, 'Checked by: ____________________'])
    ws.append(['Item ID', 'Category', 'Description', 'Count Unit', 'Physical Count', 'Notes'])
    header_fill = PatternFill('solid', fgColor='DDDDDD')
    cat_fill = PatternFill('solid', fgColor='EEEEEE')
    thin = Side(style='thin', color='333333')
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[3]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = box
    current_category = object()
    for item in _count_items(take, show_all, categories):
        if item.category != current_category:
            current_category = item.category
            ws.append([None, None, (item.category or 'Uncategorised').upper()])
            for cell in ws[ws.max_row]:
                cell.fill = cat_fill
                cell.font = Font(bold=True)
                cell.border = box
        counted = counts.get(item.id)
        ws.append([item.id, item.category or '', item.name, item.count_unit_label(),
                   float(counted) if counted is not None else None, None])
        for cell in ws[ws.max_row]:
            cell.border = box
        ws.cell(ws.max_row, 1).alignment = Alignment(horizontal='center')
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].hidden = True
    ws.column_dimensions['C'].width = 42
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 22
    ws.freeze_panes = 'A4'
    ws.print_title_rows = '3:3'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = 'Page &P of &N'
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    suffix = ('_' + secure_filename('-'.join(c or 'Uncategorised' for c in categories)[:40])
              if categories else '')
    name = f'stock_take_{take.location.name.replace(" ", "_")}_{take.count_date:%Y-%m-%d}{suffix}.xlsx'
    return send_file(buffer, as_attachment=True, download_name=name,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@inventory_bp.route('/stocktakes/<int:take_id>/variance')
def variance(take_id):
    take = StockTake.query.get_or_404(take_id)
    summary = period_summary(take)
    return render_template('inventory/stocktake_variance.html', take=take, summary=summary,
                           print_mode=False)


@inventory_bp.route('/stocktakes/<int:take_id>/print')
def variance_print(take_id):
    take = StockTake.query.get_or_404(take_id)
    summary = period_summary(take)
    return render_template('inventory/variance_print.html', take=take, summary=summary)


@inventory_bp.route('/stocktakes/<int:take_id>/reverse', methods=['POST'])
def reverse_stocktake(take_id):
    take = StockTake.query.get_or_404(take_id)
    later = StockTake.query.filter(
        StockTake.location_id == take.location_id, StockTake.status == 'confirmed',
        StockTake.count_date > take.count_date,
    ).first()
    if later:
        flash('A later stock take for this location is confirmed — reverse that one first', 'error')
        return redirect(url_for('inventory.variance', take_id=take.id))
    StockMovement.query.filter_by(stock_take_id=take.id).delete()
    for line in take.lines:
        line.expected_base = None
        line.unit_cost = None
    take.status = 'draft'
    take.confirmed_at = None
    db.session.commit()
    flash('Stock take reversed — counts are editable again', 'success')
    return redirect(url_for('inventory.count_stocktake', take_id=take.id))


@inventory_bp.route('/stocktakes/<int:take_id>/delete', methods=['POST'])
def delete_stocktake(take_id):
    take = StockTake.query.get_or_404(take_id)
    if take.status == 'confirmed':
        flash('Reverse the stock take before deleting it', 'error')
        return redirect(url_for('inventory.variance', take_id=take.id))
    db.session.delete(take)
    db.session.commit()
    flash('Stock take deleted', 'success')
    return redirect(url_for('inventory.stocktakes'))


# ---------------------------------------------------------------------------
# Testing tools
# ---------------------------------------------------------------------------

def _backup_database() -> Path | None:
    """Copy the SQLite database to instance/backups before a destructive reset."""
    import sqlite3
    db_path = db.engine.url.database
    if not db_path or db_path == ':memory:':
        return None
    backup_dir = Path(current_app.instance_path) / 'backups'
    backup_dir.mkdir(parents=True, exist_ok=True)
    target = backup_dir / f'cashup-before-reset-{datetime.now():%Y%m%d-%H%M%S}.db'
    db.session.commit()
    src = sqlite3.connect(db_path)
    dst = sqlite3.connect(str(target))
    try:
        src.backup(dst)
    finally:
        dst.close()
        src.close()
    return target


def _upload_created_items(lines, created_after) -> set[int]:
    """Items first created by these uploads (made at/after the upload started)."""
    ids = set()
    for line, started in zip(lines, created_after):
        item = CostItem.query.get(line.item_id) if line.item_id else None
        if item is not None and item.created_at >= started:
            ids.add(item.id)
    return ids


def _item_still_used(item_id) -> bool:
    return any((
        RecipeLine.query.filter_by(item_id=item_id).first(),
        StockMovement.query.filter_by(item_id=item_id).first(),
        RequisitionLine.query.filter_by(item_id=item_id).first(),
        StockTakeLine.query.filter_by(item_id=item_id).first(),
        InvoiceLine.query.filter_by(item_id=item_id).first(),
        SalesLine.query.filter_by(item_id=item_id).first(),
    ))


@inventory_bp.route('/testing', methods=['GET', 'POST'])
def testing_tools():
    counts = {
        'sales': SalesImport.query.count(),
        'invoices': Invoice.query.count(),
        'pos_aliases': PosItemAlias.query.count(),
        'supplier_aliases': SupplierItemAlias.query.count(),
        'categorised': CostItem.query.filter(CostItem.category.isnot(None),
                                             CostItem.category != '').count(),
    }
    if request.method == 'GET':
        return render_template('inventory/testing.html', counts=counts)

    form = request.form
    if (form.get('confirm_text') or '').strip().upper() != 'CLEAR':
        flash('Type CLEAR in the box to confirm', 'error')
        return redirect(url_for('inventory.testing_tools'))
    clear_sales = form.get('clear_sales') == 'on'
    clear_invoices = form.get('clear_invoices') == 'on'
    clear_items = form.get('clear_items') == 'on'
    clear_categories = form.get('clear_categories') == 'on'
    if not any((clear_sales, clear_invoices, clear_categories)):
        flash('Tick at least one thing to clear', 'error')
        return redirect(url_for('inventory.testing_tools'))

    backup = _backup_database()
    done = []
    candidate_items: set[int] = set()
    empty_recipe_ids: set[int] = set()

    if clear_sales:
        imports = SalesImport.query.all()
        lines = [l for s in imports for l in s.lines]
        candidate_items |= _upload_created_items(lines, [s.created_at for s in imports
                                                         for _ in s.lines])
        empty_recipe_ids = {l.recipe_id for l in lines if l.recipe_id}
        for s in imports:
            StockMovement.query.filter_by(sales_import_id=s.id).delete()
            db.session.delete(s)
        n_alias = PosItemAlias.query.delete()
        db.session.flush()
        for recipe_id in empty_recipe_ids:
            recipe = Recipe.query.get(recipe_id)
            if (recipe is not None and not recipe.lines
                    and (recipe.notes or '').startswith('Created from sales mix')):
                db.session.delete(recipe)
        done.append(f'{len(imports)} sales mix upload(s) and {n_alias} remembered POS mapping(s)')

    if clear_invoices:
        invoices = Invoice.query.all()
        lines = [l for inv in invoices for l in inv.lines]
        candidate_items |= _upload_created_items(lines, [inv.created_at for inv in invoices
                                                         for _ in inv.lines])
        for inv in invoices:
            StockMovement.query.filter_by(invoice_id=inv.id).delete()
            if inv.stored_filename:
                path = _upload_dir('invoices') / inv.stored_filename
                if path.exists():
                    path.unlink()
            db.session.delete(inv)
        n_alias = SupplierItemAlias.query.delete()
        done.append(f'{len(invoices)} invoice(s) and {n_alias} remembered supplier mapping(s)')

    db.session.flush()
    if clear_items and candidate_items:
        removed = 0
        for item_id in candidate_items:
            if not _item_still_used(item_id):
                db.session.delete(CostItem.query.get(item_id))
                removed += 1
        done.append(f'{removed} item(s) that those uploads had created')

    if clear_categories:
        n = CostItem.query.update({CostItem.category: None}, synchronize_session=False)
        done.append(f'categories on {n} item(s)')

    db.session.commit()
    flash('Cleared ' + '; '.join(done) + '.'
          + (f' Backup saved to {backup}' if backup else ''), 'success')
    return redirect(url_for('inventory.testing_tools'))


@inventory_bp.app_context_processor
def inject_inventory_helpers():
    return {'movement_kinds': MOVEMENT_KINDS}

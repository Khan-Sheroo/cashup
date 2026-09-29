"""Import items and recipes from the 'Recipe Costing' Excel workbook.

Workbook layout:
  * 'Item Master': Item Group | Item Name | Purchase Qty | Pur. Unit | AP Price Excl. tax |
    Base Unit (Kilogram/Liter/Each) | Based Price | Yield % | Yield Price | Notes
  * One sheet per recipe: name in A1, selling price (incl. tax) in G2, yield portions in D8,
    ingredient rows from row 12 (Ingredient | Recipe Qty | Unit | Ing. Weight | Yield % | Act. Qty | Cost),
    'Cost Add-ons' rows after the ingredients, 'Notes' further down.
  * Sheets whose name contains 'Internal' are sub-recipes (sauces, breads). The app has no nested
    recipes, so they are exploded into their raw ingredients wherever a dish uses them.
"""
from __future__ import annotations

import difflib
import io
import re
from dataclasses import dataclass, field
from decimal import Decimal

from cashup import db
from cashup.category_rules import guess_category
from cashup.models import CostItem, Location, PosItemAlias, Recipe, RecipeLine
from cashup.recipe_util import unit_conversion

SKIP_SHEETS = {'measurement', 'recipe template', 'item master'}
BASE_UNITS = {'kilogram': ('g', Decimal('1000'), 'kg'), 'liter': ('ml', Decimal('1000'), 'L'),
              'litre': ('ml', Decimal('1000'), 'L'), 'each': ('each', Decimal('1'), None)}
PURCHASE_TO_BASE = {'kg': Decimal('1'), 'g': Decimal('0.001'), 'l': Decimal('1'), 'ml': Decimal('0.001'),
                    'no': Decimal('1'), 'each': Decimal('1'), 'pcs': Decimal('1')}
GROUP_CATEGORIES = {
    'meat': 'Meat', 'fish': 'Seafood', 'vegetable': 'Vegetables', 'fruits': 'Fruit', 'fruit': 'Fruit',
    'dairy': 'Dairy & Eggs', 'juices': 'Juices', 'spices': 'Spices', 'sauce': 'Sauces & Condiments',
    'oil': 'Oils & Vinegars', 'bakery': 'Bakery', 'frozan': 'Frozen', 'frozen': 'Frozen',
    'ice cream': 'Desserts', 'puree': 'Juices', 'dry': 'Dry Goods',
}
ZERO = Decimal('0')


def _key(name) -> str:
    return re.sub(r'\s+', ' ', str(name or '')).strip().lower()


def _dec(value) -> Decimal | None:
    if value in (None, ''):
        return None
    try:
        return Decimal(str(value))
    except Exception:  # noqa: BLE001
        return None


@dataclass
class MasterItem:
    name: str
    group: str
    base_unit: str          # 'Kilogram' / 'Liter' / 'Each'
    purchase_base: Decimal  # purchase quantity in base units (kg / L / each)
    price: Decimal          # AP price excl. tax for that purchase quantity
    notes: str


@dataclass
class SheetRecipe:
    sheet: str
    name: str
    price: Decimal | None
    portions: Decimal
    lines: list[tuple[str, Decimal]]  # (ingredient name, amount in the ingredient's base unit per batch)
    notes: list[str]
    portions_given: bool = True


@dataclass
class ImportReport:
    items_created: list[str] = field(default_factory=list)
    items_reused: list[str] = field(default_factory=list)
    recipes_created: list[str] = field(default_factory=list)
    recipes_updated: list[str] = field(default_factory=list)
    pos_linked: list[tuple[str, str]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def read_workbook(file_or_path):
    import warnings

    import openpyxl
    if hasattr(file_or_path, 'read'):
        file_or_path = io.BytesIO(file_or_path.read())
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        wb = openpyxl.load_workbook(file_or_path, data_only=True, read_only=True)

    master: dict[str, MasterItem] = {}
    for row in list(wb['Item Master'].iter_rows(values_only=True))[2:]:
        if len(row) < 6 or not row[1]:
            continue
        base = str(row[5] or '').strip().lower()
        qty, price = _dec(row[2]), _dec(row[4])
        factor = PURCHASE_TO_BASE.get(str(row[3] or '').strip().lower())
        if base not in BASE_UNITS or not qty or price is None or factor is None:
            continue
        name = re.sub(r'\s+', ' ', str(row[1])).strip()
        master[_key(name)] = MasterItem(name, str(row[0] or '').strip(), base, qty * factor, price,
                                        str(row[9] or '').strip() if len(row) > 9 else '')

    recipes: list[SheetRecipe] = []
    for ws in wb.worksheets:
        if ws.title.strip().lower() in SKIP_SHEETS:
            continue
        rows = [list(r) + [None] * 14 for r in ws.iter_rows(values_only=True)]
        if len(rows) < 12:
            continue
        lines = []
        section_end = len(rows)
        for idx in range(11, len(rows)):
            first = rows[idx][0]
            if isinstance(first, str) and first.strip().lower() == 'notes':
                section_end = idx
                break
        for row in rows[11:section_end]:
            name = row[0]
            if not isinstance(name, str) or not name.strip():
                continue
            if name.strip().lower() in ('main recipe', 'cost add-ons'):
                continue
            weight = _dec(row[3])
            yield_pct = _dec(row[4]) or Decimal('1')
            actual = _dec(row[5])
            if actual is None and weight is not None:
                actual = weight / yield_pct
            if not actual:
                continue
            lines.append((re.sub(r'\s+', ' ', name).strip(), actual))
        if not lines:
            continue
        notes = [str(r[0]).strip() for r in rows[section_end + 1:] if r[0]]
        recipes.append(SheetRecipe(
            sheet=ws.title,
            name=re.sub(r'\s+', ' ', str(rows[0][0] or ws.title)).strip(),
            price=_dec(rows[1][6]),
            portions=_dec(rows[7][3]) or Decimal('1'),
            lines=lines,
            notes=notes,
            portions_given=_dec(rows[7][3]) is not None,
        ))
    return master, recipes


def _is_internal(name: str) -> bool:
    return 'internal' in name.lower()


def _match(name: str, keys) -> str | None:
    key = _key(name)
    if key in keys:
        return key
    close = difflib.get_close_matches(key, list(keys), n=1, cutoff=0.88)
    return close[0] if close else None


def _explode(recipe: SheetRecipe, master, internals, report, depth=0) -> list[tuple[str, Decimal]]:
    """Batch lines with Internal sub-recipes replaced by their raw ingredients."""
    out = []
    for name, amount in recipe.lines:
        if _is_internal(name) and depth < 3:
            sub_key = _match(name, internals)
            if sub_key is not None:
                sub = internals[sub_key]
                batch = _internal_batch_size(sub, master)
                if batch:
                    for sub_name, sub_amount in _explode(sub, master, internals, report, depth + 1):
                        out.append((sub_name, sub_amount * amount / batch))
                    continue
                report.warnings.append(f'{recipe.name}: could not size sub-recipe "{name}" — kept as an item')
        out.append((name, amount))
    return out


def _internal_batch_size(sub: SheetRecipe, master) -> Decimal | None:
    """How many base units (kg / L / each) one batch of an Internal sheet makes."""
    info = master.get(_key(sub.name))
    if info and info.base_unit == 'each':
        return sub.portions or Decimal('1')
    if info and info.purchase_base:
        return info.purchase_base
    total = ZERO
    for name, amount in sub.lines:
        m = master.get(_match(name, master) or '')
        if m and m.base_unit != 'each':
            total += amount
    return total or None


def _existing_item_for(m: MasterItem):
    """An existing stock item with the same name and a compatible unit (sales mix items excluded)."""
    unit, per_base, _count = BASE_UNITS[m.base_unit]
    names = [f'{m.name} (Kitchen)'.lower(), m.name.lower()]
    candidates = CostItem.query.filter(db.func.lower(CostItem.name).in_(names)).all()
    candidates.sort(key=lambda i: names.index(i.name.lower()))
    for item in candidates:
        if item.category == 'From Sales Mix':
            continue
        conv = Decimal('1') if unit == item.unit else unit_conversion(unit, item.unit)
        if conv is None and unit != 'each' and item.unit in ('g', 'kg', 'ml', 'L', 'l'):
            conv = unit_conversion('g' if unit == 'ml' else 'ml', item.unit)  # treat ml ≈ g like the workbook
        if conv is not None:
            return item, per_base * conv
    return None, None


def _item_for_master(m: MasterItem, report: ImportReport, cache: dict):
    """(CostItem, recipe units per base unit) for an Item Master row, creating it when needed."""
    key = _key(m.name)
    if key in cache:
        return cache[key]
    item, factor = _existing_item_for(m)
    if item is not None:
        report.items_reused.append(item.name)
    else:
        unit, per_base, count_unit = BASE_UNITS[m.base_unit]
        name = m.name
        clash = CostItem.query.filter(db.func.lower(CostItem.name) == name.lower()).first()
        if clash is not None and clash.category != 'From Sales Mix':
            name = f'{m.name} (Kitchen)'
        category = guess_category(m.name) or GROUP_CATEGORIES.get(m.group.strip().lower())
        if m.name.lower().startswith('other cost'):
            category = 'Other Costs'
        existing = next((c for (c,) in db.session.query(CostItem.category).distinct()
                         if c and category and c.lower() == category.lower()), None)
        item = CostItem(
            name=name[:200], unit=unit, pack_size=m.purchase_base * per_base, cost_price=m.price,
            category=existing or category, count_unit=count_unit,
            count_factor=per_base if count_unit else None, active=True,
        )
        db.session.add(item)
        db.session.flush()
        factor = per_base
        report.items_created.append(item.name)
    cache[key] = (item, factor)
    return cache[key]


def import_workbook(file_or_path, source_name='Recipe Costing workbook',
                    import_all_items=True, link_pos=True) -> ImportReport:
    master, sheets = read_workbook(file_or_path)
    report = ImportReport()
    internals = {_key(s.name): s for s in sheets if _is_internal(s.name)}
    kitchen = Location.query.filter(db.func.lower(Location.name) == 'kitchen').first()
    cache: dict = {}

    if import_all_items:
        for key, m in master.items():
            if not _is_internal(m.name):
                _item_for_master(m, report, cache)

    for sheet in sheets:
        if _is_internal(sheet.name):
            continue
        usage: dict[int, Decimal] = {}
        order: list[int] = []
        for name, amount in _explode(sheet, master, internals, report):
            mkey = _match(name, master)
            if mkey is None:
                report.warnings.append(f'{sheet.name}: "{name}" is not in the Item Master — skipped')
                continue
            item, factor = _item_for_master(master[mkey], report, cache)
            qty = amount * factor / sheet.portions
            if item.id not in usage:
                order.append(item.id)
            usage[item.id] = usage.get(item.id, ZERO) + qty

        recipe = Recipe.query.filter(db.func.lower(Recipe.name) == sheet.name.lower()).first()
        note_lines = [f'Imported from {source_name} (sheet "{sheet.sheet}").']
        if not sheet.portions_given:
            note_lines.append('No yield set in the workbook: quantities are for the WHOLE BATCH.')
            report.warnings.append(f'{sheet.name}: no yield in the workbook, imported as one whole batch')
        elif sheet.portions != 1:
            note_lines.append(f'Batch recipe: yields {sheet.portions.normalize():f} portions; '
                              f'quantities here are per portion.')
        note_lines += sheet.notes
        if recipe is None:
            recipe = Recipe(name=sheet.name[:200])
            db.session.add(recipe)
            report.recipes_created.append(sheet.name)
        else:
            recipe.lines.clear()
            db.session.flush()
            report.recipes_updated.append(sheet.name)
        recipe.sale_price = sheet.price.quantize(Decimal('0.01')) if sheet.price else recipe.sale_price
        recipe.notes = '\n'.join(note_lines)
        if kitchen is not None and recipe.location_id is None:
            recipe.location_id = kitchen.id
        for sort, item_id in enumerate(order):
            recipe.lines.append(RecipeLine(item_id=item_id, quantity=usage[item_id].quantize(Decimal('0.0001')),
                                           sort_order=sort))
        db.session.flush()

    if link_pos:
        _link_pos_names(report, [s.name for s in sheets if not _is_internal(s.name)])
    return report


def _pos_key(name: str) -> str:
    text = _key(name)
    text = re.sub(r'\(.*?\)|\b(the|for|kids?|side|portion)\b|[^a-z0-9 ]', ' ', text)
    return re.sub(r'\s+', ' ', text).strip()


def _link_pos_names(report: ImportReport, recipe_names: list[str]):
    """Point remembered POS names that are sold as plain items at the matching new recipe."""
    targets = {}
    for name in recipe_names:
        recipe = Recipe.query.filter(db.func.lower(Recipe.name) == name.lower()).first()
        if recipe is not None:
            targets[name] = recipe
    for alias in PosItemAlias.query.filter(PosItemAlias.recipe_id.is_(None)).all():
        item = CostItem.query.get(alias.item_id) if alias.item_id else None
        if alias.ignore or (item is not None and item.category != 'From Sales Mix'):
            continue
        pos_name = item.name if item is not None else alias.pos_name_key
        pos_is_kids = bool(re.search(r'\bkids?\b', pos_name, re.I))
        best, best_score = None, 0.0
        for name, recipe in targets.items():
            if bool(re.search(r'\bkids?\b', name, re.I)) != pos_is_kids:
                continue
            score = difflib.SequenceMatcher(None, _pos_key(pos_name), _pos_key(name)).ratio()
            if score > best_score:
                best, best_score = recipe, score
        if best is not None and best_score >= 0.8:
            alias.recipe_id = best.id
            alias.item_id = None
            report.pos_linked.append((pos_name, best.name))
